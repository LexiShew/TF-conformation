#!/usr/bin/env python
"""Assemble the v5 TF-conformation talk on the August 2026 deck's own template.

Why it works this way: python-pptx cannot copy a slide from one file into
another, so the only way to keep BOTH the house theme (Compact Stripe / Ocean
Fairy, Aptos, 10 x 5.63 in) AND a set of existing slides verbatim is to start
from a COPY of the August deck, delete the slides we are not reusing, append the
new ones on its own layouts, and then reorder.

  carry   -- an August slide kept byte-for-byte (speaker notes may be replaced)
  title   -- the opening slide, on the template's "Title Slide" layout
  section -- a divider, on "Section Header"
  text    -- title + bullets, on "Title (Short) and Content"
  figure  -- title + one or two images (+ optional caption), on "Title (Short)"

Speaker notes are bullet lines, per the talk plan.

Usage
-----
    python build_v5_deck.py \
        --template 8-06-2026_TF_conformations.pptx \
        --spec     deck_spec_v5.json \
        --figures  deckimg \
        --out      TF_conformation_v5_deck.pptx

`--figures` holds the images named in the spec. Regenerate the N1-N10 panels
with analysis/analyses/nmr_rmsd/make_talk_figs.py; the rest are existing repo
figures (see FIGURE_SOURCES in deck_spec_v5.json).

Deleting slides drops their relationships, so the media they carried is not
serialized either -- the output is a few MB rather than the template's 77.
"""
import argparse
import json
import os

from PIL import Image
from pptx import Presentation
from pptx.opc.packuri import PackURI
from pptx.util import Inches, Pt

RELID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"

# Template geometry, measured from the August deck (inches).
SLIDE_W, SLIDE_H = 10.0, 5.63
TITLE = (0.25, 0.16, 9.5, 0.66)          # a little taller than the layout box so
                                         # two-line titles do not clip
BODY = (0.32, 0.92, 9.36, 4.40)
FIG_FULL = (0.30, 0.86, 9.40, 4.46)
FIG_CAPPED = (0.30, 0.84, 9.40, 4.00)
CAPTION = (0.32, 4.92, 9.36, 0.58)


# --------------------------------------------------------------------------- #
# slide-list surgery
# --------------------------------------------------------------------------- #
def drop_slide(prs, element):
    """Remove one slide by its sldId element, dropping the rel so the slide part
    (and any media only it referenced) is left out of the saved package."""
    prs.part.drop_rel(element.get(RELID))
    prs.slides._sldIdLst.remove(element)


def keep_only(prs, keep_1based):
    """Delete every slide except the given 1-based positions. Returns the kept
    sldId elements in their original order."""
    elements = list(prs.slides._sldIdLst)
    keep = {i - 1 for i in keep_1based}
    kept = [e for i, e in enumerate(elements) if i in keep]
    for i, e in enumerate(elements):
        if i not in keep:
            drop_slide(prs, e)
    return kept


def renumber(prs, kept):
    """Rename the surviving slide parts to slide1..slideN.

    Needed because add_slide() names a new part /ppt/slides/slide<len+1>.xml
    while carried slides keep their ORIGINAL names (slide29.xml, slide103.xml,
    ...). Without this, a new slide can be handed a name a carried slide already
    holds and the two collide in the saved zip ("Duplicate name:
    ppt/slides/slide29.xml"), which loses a slide. Renaming in ascending
    original order is collision-free: any carried part whose number falls in the
    target range is itself renamed before that number is needed.
    Relationships are rebuilt from partnames at save time, so nothing else
    needs updating.
    """
    for i, sldId in enumerate(kept, start=1):
        part = prs.part.related_part(sldId.get(RELID))
        part.partname = PackURI(f"/ppt/slides/slide{i}.xml")


def reorder(prs, elements):
    """Rewrite the slide order to exactly `elements`."""
    lst = prs.slides._sldIdLst
    current = list(lst)
    assert len(current) == len(elements), f"{len(current)} slides vs {len(elements)} ordered"
    for e in current:
        lst.remove(e)
    for e in elements:
        lst.append(e)


# --------------------------------------------------------------------------- #
# slide construction
# --------------------------------------------------------------------------- #
def layout(prs, name):
    for lo in prs.slide_masters[0].slide_layouts:
        if lo.name == name:
            return lo
    raise KeyError(f"layout {name!r} not in template: "
                   f"{[lo.name for lo in prs.slide_masters[0].slide_layouts]}")


def set_notes(slide, lines):
    if not lines:
        return
    slide.notes_slide.notes_text_frame.text = "\n".join("\u2022  " + x for x in lines)


def put_title(slide, text):
    """Fill the layout's title placeholder, nudged taller for two-line titles.
    Font name is left alone so the theme's Aptos Display is inherited."""
    ph = slide.shapes.title
    ph.left, ph.top = Inches(TITLE[0]), Inches(TITLE[1])
    ph.width, ph.height = Inches(TITLE[2]), Inches(TITLE[3])
    tf = ph.text_frame
    tf.word_wrap = True
    tf.text = text
    size = 20 if len(text) <= 62 else (18 if len(text) <= 92 else 16)
    for p in tf.paragraphs:
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.bold = True
    return ph


def put_bullets(slide, bullets):
    """Bullets into the layout's body placeholder (keeps theme list formatting).

    A string is a top-level bullet; a [string] of length 1 nested one level is
    written as a sub-bullet, i.e. spec entries may be "text" or ["text"].
    """
    body = None
    for ph in slide.placeholders:
        if ph.placeholder_format.idx != 0 and ph.has_text_frame:
            body = ph
            break
    assert body is not None, "layout has no body placeholder"
    body.left, body.top = Inches(BODY[0]), Inches(BODY[1])
    body.width, body.height = Inches(BODY[2]), Inches(BODY[3])
    tf = body.text_frame
    tf.word_wrap = True
    tf.clear()
    first = True
    for b in bullets:
        sub = isinstance(b, list)
        text = b[0] if sub else b
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = 1 if sub else 0
        run = p.add_run()
        run.text = text
        run.font.size = Pt(12 if sub else 14)
        p.space_after = Pt(5 if sub else 8)
    return body


def place_image(slide, path, box):
    """Fit an image inside `box` (inches), centred, preserving aspect ratio."""
    x, y, bw, bh = box
    iw, ih = Image.open(path).size
    scale = min(bw / (iw / 96), bh / (ih / 96))      # PIL px -> in at 96 dpi
    w, h = (iw / 96) * scale, (ih / 96) * scale
    return slide.shapes.add_picture(
        path, Inches(x + (bw - w) / 2), Inches(y + (bh - h) / 2),
        Inches(w), Inches(h))


def put_caption(slide, text):
    box = slide.shapes.add_textbox(*[Inches(v) for v in CAPTION])
    tf = box.text_frame
    tf.word_wrap = True
    tf.text = text
    for p in tf.paragraphs:
        for r in p.runs:
            r.font.size = Pt(10)
            r.font.italic = True
    return box


# --------------------------------------------------------------------------- #
def build(template, spec, figdir, out):
    prs = Presentation(template)
    assert len(prs.slides._sldIdLst) > 40, "template does not look like the August deck"

    carried = [s["old"] for s in spec if s["kind"] == "carry"]
    assert len(carried) == len(set(carried)), "an August slide is carried twice"
    kept = keep_only(prs, carried)
    kept_by_old = dict(zip(sorted(carried), kept))   # deletion preserves order
    renumber(prs, kept)                              # before any add_slide()

    order = []
    for item in spec:
        kind = item["kind"]
        if kind == "carry":
            element = kept_by_old[item["old"]]
            idx = list(prs.slides._sldIdLst).index(element)
            set_notes(prs.slides[idx], item.get("notes"))
            order.append(element)
            continue

        if kind == "title":
            slide = prs.slides.add_slide(layout(prs, "Title Slide"))
            tf = slide.shapes.title.text_frame
            tf.word_wrap = True
            tf.text = item["title"]
            for p in tf.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(26)
                    r.font.bold = True
            subs = [ph for ph in slide.placeholders if ph.placeholder_format.idx != 0
                    and ph.has_text_frame]
            if subs and item.get("subtitle"):
                subs[0].text_frame.word_wrap = True
                subs[0].text_frame.text = item["subtitle"]
                for p in subs[0].text_frame.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(13)
        elif kind == "section":
            slide = prs.slides.add_slide(layout(prs, "Section Header"))
            tf = slide.shapes.title.text_frame
            tf.word_wrap = True
            tf.text = item["title"]
            for p in tf.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(24)
                    r.font.bold = True
            for ph in slide.placeholders:
                if ph.placeholder_format.idx != 0 and ph.has_text_frame and item.get("subtitle"):
                    ph.text_frame.text = item["subtitle"]
                    for p in ph.text_frame.paragraphs:
                        for r in p.runs:
                            r.font.size = Pt(12)
        elif kind == "text":
            slide = prs.slides.add_slide(layout(prs, "Title (Short) and Content"))
            put_title(slide, item["title"])
            put_bullets(slide, item["bullets"])
        elif kind == "figure":
            slide = prs.slides.add_slide(layout(prs, "Title (Short)"))
            put_title(slide, item["title"])
            images = [os.path.join(figdir, f) for f in item["images"]]
            for p in images:
                assert os.path.exists(p), f"missing figure: {p}"
            box = FIG_CAPPED if item.get("caption") else FIG_FULL
            if len(images) == 1:
                place_image(slide, images[0], box)
            else:
                # Two wide panels side by side shrink to nothing (each 4.6 in
                # wide -> 1.9 in tall at aspect 2.4). Stack wide images in rows
                # and only put narrow ones in columns.
                x, y, bw, bh = box
                gutter, n = 0.16, len(images)
                aspects = [Image.open(q).size[0] / Image.open(q).size[1] for q in images]
                if min(aspects) > 1.8:
                    h = (bh - gutter * (n - 1)) / n
                    for i, q in enumerate(images):
                        place_image(slide, q, (x, y + i * (h + gutter), bw, h))
                else:
                    w = (bw - gutter * (n - 1)) / n
                    for i, q in enumerate(images):
                        place_image(slide, q, (x + i * (w + gutter), y, w, bh))
            if item.get("caption"):
                put_caption(slide, item["caption"])
        else:
            raise ValueError(f"unknown slide kind {kind!r}")

        set_notes(slide, item.get("notes"))
        order.append(list(prs.slides._sldIdLst)[-1])

    reorder(prs, order)
    prs.save(out)
    return prs


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--template", required=True, help="the August 2026 pptx")
    ap.add_argument("--spec", default="deck_spec_v5.json")
    ap.add_argument("--figures", default="deckimg")
    ap.add_argument("--out", default="TF_conformation_v5_deck.pptx")
    a = ap.parse_args()
    spec = [s for s in json.load(open(a.spec)) if not s.get("_comment")]
    prs = build(a.template, spec, a.figures, a.out)
    print("wrote %s -- %d slides, %.1f MB"
          % (a.out, len(prs.slides._sldIdLst), os.path.getsize(a.out) / 1e6))


if __name__ == "__main__":
    main()
